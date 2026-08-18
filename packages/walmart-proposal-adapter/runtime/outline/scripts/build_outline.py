#!/usr/bin/env python3
"""Build a deterministic Walmart buyer proposal outline from analysis artifacts."""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
from collections import defaultdict
from datetime import datetime, timezone
from difflib import get_close_matches
from pathlib import Path
from typing import Any

from artifact_inputs import load_manifest


TBD_FIELDS = [
    "sku", "upc", "product_name", "cost", "retail_price", "margin",
    "sales_forecast", "store_count", "launch_date", "pack_size",
    "walmart_commitment",
]
POSITIVE_WS = {
    "expand_distribution": 1.0,
    "qualified_entry_opportunity": 0.95,
    "entry_opportunity": 0.95,
    "defend_or_deepen": 0.85,
    "improve_or_rationalize": 0.70,
    "test_and_learn": 0.65,
    "avoid_crowded_space": 0.20,
}
OPPORTUNITY_SCORE = {
    "new_product_opportunity": 1.0,
    "space_expansion_candidate": 0.9,
    "refresh_or_repack_opportunity": 0.8,
    "supplier_validation_required": 0.62,
    "watchlist": 0.55,
    "sku_rationalization_candidate": 0.45,
    "deprioritize": 0.15,
}
PRODUCTIVITY_SCORE = {
    "productivity_improving": 1.0,
    "matched_sku_growth": 0.9,
    "sku_expansion_led": 0.55,
    "replacement_growth": 0.55,
    "declining_existing_skus": 0.35,
    "insufficient_sku_data": 0.25,
    "insufficient_match_data": 0.25,
}
CONCENTRATION_SCORE = {
    "broad_based": 1.0,
    "moderate_concentration": 0.7,
    "high_concentration": 0.35,
    "no_current_sales": 0.1,
}
INVESTMENT_ELIGIBLE_TYPES = {
    "new_product_opportunity",
    "refresh_or_repack_opportunity",
    "space_expansion_candidate",
    "sku_rationalization_candidate",
}


def load_json(path: Path) -> dict[str, Any]:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        raise SystemExit(f"Cannot read valid JSON from {path}: {exc}") from exc


def norm(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", " ", str(value or "").upper()).strip()


def category_key(value: Any) -> str:
    """Normalize harmless category naming variants such as CRAFT/CRAFTS."""
    tokens = norm(value).split()
    return " ".join(token[:-1] if token.endswith("S") and len(token) > 3 else token for token in tokens)


def slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def percentile(values: list[float], value: float) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    if len(ordered) == 1:
        return 1.0
    below = sum(v < value for v in ordered)
    equal = sum(v == value for v in ordered)
    return (below + max(equal - 1, 0) / 2) / (len(ordered) - 1)


def finite(value: Any, default: float = 0.0) -> float:
    try:
        number = float(value)
        return number if math.isfinite(number) else default
    except (TypeError, ValueError):
        return default


def money(value: Any) -> str:
    number = finite(value)
    if abs(number) >= 1_000_000:
        return f"${number / 1_000_000:.1f}M"
    if abs(number) >= 1_000:
        return f"${number / 1_000:.1f}K"
    return f"${number:,.0f}"


def pct(value: Any, digits: int = 1) -> str:
    if value is None:
        return "N/A"
    return f"{finite(value) * 100:+.{digits}f}%"


def plain_pct(value: Any, digits: int = 1) -> str:
    if value is None:
        return "N/A"
    return f"{finite(value) * 100:.{digits}f}%"


def direction(value: Any) -> str:
    number = finite(value)
    return "up" if number > 0 else "down" if number < 0 else "flat"


class PathFormatter:
    def __init__(self, mode: str, base_dir: Path | None = None) -> None:
        if mode not in {"relative", "absolute"}:
            raise ValueError("path mode must be relative or absolute")
        self.mode = mode
        self.base_dir = (base_dir or Path.cwd()).resolve()

    def format(self, path: Path | None) -> str | None:
        if path is None:
            return None
        resolved = path.resolve()
        if self.mode == "absolute":
            return str(resolved)
        try:
            return resolved.relative_to(self.base_dir).as_posix()
        except ValueError:
            return str(path)


class Evidence:
    def __init__(self, paths: PathFormatter) -> None:
        self.records: list[dict[str, Any]] = []
        self.counter = 0
        self.paths = paths

    def add(self, source_file: Path, json_path: str, value: Any, statement_type: str,
            period: str | None = None, unit: str | None = None,
            note: str | None = None) -> str:
        self.counter += 1
        evidence_id = f"PE{self.counter:05d}"
        source_name = source_file.name.lower()
        source_layer = (
            "white_space_analysis" if "white_space" in source_name
            else "trend_analysis" if "analysis_package" in source_name
            else "fineline_investment_analysis" if "fineline" in source_name
            else "product_input"
        )
        self.records.append({
            "evidence_id": evidence_id,
            "source_file": self.paths.format(source_file),
            "source_json_path": json_path,
            "value": value,
            "unit": unit,
            "period": period,
            "statement_type": statement_type,
            "source_layer": source_layer,
            "note": note,
        })
        return evidence_id


def card(label: str, raw: Any, display: str, unit: str, period: str,
         comparison: str, interpretation: str, source: Path, json_path: str,
         ev: Evidence, direction_value: Any | None = None) -> dict[str, Any]:
    eid = ev.add(source, json_path, raw, "fact", period, unit, interpretation)
    return {
        "label": label, "display_value": display, "raw_value": raw, "unit": unit,
        "period": period, "comparison_basis": comparison,
        "direction": direction(direction_value if direction_value is not None else raw),
        "interpretation": interpretation, "source_file": ev.paths.format(source),
        "source_json_path": json_path, "evidence_ids": [eid],
    }


def insight(text: str, statement_type: str, evidence_ids: list[str]) -> dict[str, Any]:
    return {"text": text, "statement_type": statement_type, "evidence_ids": evidence_ids}


def load_products(path: Path | None) -> list[dict[str, Any]]:
    if not path:
        return []
    if not path.exists():
        raise SystemExit(f"Product input does not exist: {path}")
    suffix = path.suffix.lower()
    if suffix == ".json":
        data = load_json(path)
        if isinstance(data, list):
            return data
        for key in ("products", "items", "records"):
            if isinstance(data.get(key), list):
                return data[key]
        return [data]
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise SystemExit("openpyxl is required for XLSX product input") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(v or "").strip() for v in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:] if any(v is not None for v in row)]
    raise SystemExit("Product input must be JSON, CSV, XLSX, or XLSM")


def image_matches(directory: Path | None, fineline: str, paths: PathFormatter) -> list[str]:
    if not directory:
        return []
    if not directory.exists():
        raise SystemExit(f"Product image directory does not exist: {directory}")
    needle = norm(fineline).replace(" ", "")
    return [paths.format(p) for p in directory.rglob("*")
            if p.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}
            and needle in norm(p.stem).replace(" ", "")]


def calculate_investment_scores(finelines: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    sales_values = [finite(row.get("primary_current_sales_value")) for row in finelines]
    growth_values = [max(-1.0, min(2.0, finite(row.get("primary_yoy_growth_pct")))) for row in finelines]
    results: dict[str, dict[str, Any]] = {}
    for row in finelines:
        name = norm(row.get("final_fineline"))
        scale = percentile(sales_values, finite(row.get("primary_current_sales_value")))
        growth = percentile(growth_values, max(-1.0, min(2.0, finite(row.get("primary_yoy_growth_pct")))))
        productivity = PRODUCTIVITY_SCORE.get(str(row.get("sku_productivity_label")), 0.45)
        opportunity = OPPORTUNITY_SCORE.get(str(row.get("opportunity_type")), 0.4)
        concentration = CONCENTRATION_SCORE.get(str(row.get("concentration_label")), 0.45)
        investment = 100 * (0.25 * scale + 0.25 * growth + 0.20 * productivity + 0.20 * opportunity + 0.10 * concentration)

        results[name] = {
            "investment_score": round(investment, 2),
            "sales_scale_score": round(100 * scale, 2),
            "growth_momentum_score": round(100 * growth, 2),
            "sku_productivity_score": round(100 * productivity, 2),
            "opportunity_role_score": round(100 * opportunity, 2),
            "concentration_quality_score": round(100 * concentration, 2),
        }
    return results


def best_ws_rows(rows: list[dict[str, Any]], limit: int = 8) -> list[dict[str, Any]]:
    def rank(row: dict[str, Any]) -> tuple[float, float, float]:
        return (
            POSITIVE_WS.get(str(row.get("strategy_code") or ""), 0.5),
            finite(row.get("current_sales_value_percentile")),
            finite(row.get("recent_8w_productivity_percentile")),
        )
    return sorted(rows, key=rank, reverse=True)[:limit]


def make_slide(slides: list[dict[str, Any]], section: str, role: str, headline: str,
               purpose: str, takeaway: str, density: str = "medium",
               fineline: str | None = None, layout: str = "headline_content",
               cards: list[dict[str, Any]] | None = None,
               chart: dict[str, Any] | None = None,
               insights: list[dict[str, Any]] | None = None,
               visuals: list[dict[str, Any]] | None = None,
               limitations: list[str] | None = None,
               product_fields: dict[str, Any] | None = None) -> dict[str, Any]:
    slide = {
        "slide_id": f"S{len(slides)+1:03d}", "section": section,
        "fineline": fineline, "slide_role": role, "headline": headline,
        "purpose": purpose, "buyer_takeaway": takeaway, "layout": layout,
        "density": density, "metric_cards": cards or [], "chart": chart,
        "insights": insights or [], "visual_assets": visuals or [],
        "limitations": limitations or [],
    }
    if product_fields is not None:
        slide["product_fields"] = product_fields
    slide["source_refs"] = sorted({eid for c in slide["metric_cards"] for eid in c.get("evidence_ids", [])}
                                  | {eid for i in slide["insights"] for eid in i.get("evidence_ids", [])})
    if chart:
        slide["source_refs"] = sorted(set(slide["source_refs"]) | set(chart.get("evidence_ids", [])))
    slides.append(slide)
    return slide


def render_markdown(spec: dict[str, Any]) -> str:
    lines = [f"# {spec['deck_title']}", "", f"Generated: {spec['generated_at']}",
             f"Category: {spec['category']}", f"Slides: {len(spec['slides'])}",
             f"Selected Finelines: {', '.join(spec['selection']['selected_finelines'])}", ""]
    for slide in spec["slides"]:
        lines += [f"## {slide['slide_id']} — {slide['headline']}", "",
                  f"- **Section:** {slide['section']}",
                  f"- **Role:** {slide['slide_role']}",
                  f"- **Fineline:** {slide['fineline'] or 'Category / Portfolio'}",
                  f"- **Purpose:** {slide['purpose']}",
                  f"- **Buyer takeaway:** {slide['buyer_takeaway']}",
                  f"- **Layout:** `{slide['layout']}`; density `{slide['density']}`", ""]
        if slide.get("fineline_context"):
            context = slide["fineline_context"]
            lines += ["### Sequential analysis status", "",
                      f"- **Selection source:** `{context['selection_source']}`",
                      f"- **Investment rank / role:** `{context['investment_rank']}` / `{context['investment_opportunity_type']}`",
                      f"- **White Space coverage:** `{context['white_space_coverage']}`",
                      f"- **Proposal readiness:** `{context['proposal_readiness']}`", ""]
        if slide["metric_cards"]:
            lines += ["### KPI cards", ""]
            for idx, c in enumerate(slide["metric_cards"], 1):
                lines += [f"{idx}. **{c['label']}: {c['display_value']}**",
                          f"   - Raw value / unit: `{c['raw_value']}` / `{c['unit']}`",
                          f"   - Period / comparison: {c['period']} / {c['comparison_basis']}",
                          f"   - Direction / interpretation: `{c['direction']}` / {c['interpretation']}",
                          f"   - Source: `{c['source_json_path']}` ({', '.join(c['evidence_ids'])})"]
            lines.append("")
        if slide.get("chart"):
            ch = slide["chart"]
            lines += ["### Main chart", "", f"- **Type:** `{ch.get('chart_type')}`",
                      f"- **Business question:** {ch.get('business_question')}",
                      f"- **X / Y / size / color:** `{ch.get('x_axis')}` / `{ch.get('y_axis')}` / `{ch.get('size')}` / `{ch.get('color')}`",
                      f"- **Sort / filters:** `{ch.get('sort')}` / `{ch.get('filters')}`",
                      f"- **Reference lines:** `{json.dumps(ch.get('reference_lines'), ensure_ascii=False)}`",
                      f"- **Highlight:** `{ch.get('highlight')}`",
                      f"- **Annotations:** {'; '.join(ch.get('annotations') or [])}",
                      f"- **Display contract:** `{json.dumps(ch.get('display'), ensure_ascii=False)}`",
                      f"- **Data rows:** `{json.dumps(ch.get('data_rows'), ensure_ascii=False)}`",
                      f"- **Source paths:** `{', '.join(ch.get('source_json_paths') or [])}`", ""]
        if slide["insights"]:
            lines += ["### Insights", ""]
            for item in slide["insights"]:
                lines.append(f"- [{item['statement_type']}] {item['text']} ({', '.join(item['evidence_ids']) or 'no numeric evidence'})")
            lines.append("")
        if slide["visual_assets"]:
            lines += ["### Visual assets", ""]
            for visual in slide["visual_assets"]:
                lines.append(f"- `{visual.get('path')}` — {visual.get('usage')} ({visual.get('evidence_role')})")
            lines.append("")
        if slide.get("product_fields") is not None:
            lines += ["### Product fields", "", "```json",
                      json.dumps(slide["product_fields"], ensure_ascii=False, indent=2), "```", ""]
        if slide["limitations"]:
            lines += ["### Limitations / validation", ""] + [f"- {x}" for x in slide["limitations"]] + [""]
    return "\n".join(lines).rstrip() + "\n"


def numeric_tokens(value: Any) -> list[str]:
    """Return numeric tokens an LLM may repeat without creating a new claim."""
    text = json.dumps(value, ensure_ascii=False, default=str)
    return sorted(set(re.findall(r"(?<![A-Za-z])[-+]?\$?\d[\d,]*(?:\.\d+)?%?", text)))


def build_narrative_brief(spec: dict[str, Any]) -> dict[str, Any]:
    briefs = []
    priority_index = next((i for i, x in enumerate(spec["slides"]) if x["slide_role"] == "fineline_ranking"), len(spec["slides"]))
    role_guidance = {
        "cover": "Use the exact corporate-retailer-category title; keep the subtitle factual and minimal.",
        "category_snapshot": "Establish category size and performance only. Do not introduce selected Finelines or recommendations yet.",
        "category_growth_drivers": "Explain what drives total-category change in merchant language: sales, demand, and average selling price.",
        "portfolio_quadrant": "Explain the category structure before revealing the chosen investment priorities.",
        "cultural_shift_overview": "Translate consumer shifts into product-development relevance without making a Walmart demand claim.",
        "trend_visual": "Use product-development language: consumer need, making behavior, product format, design code, and usage occasion.",
        "fineline_ranking": "Introduce the investment priorities only after category and trend context have established the case.",
        "business_case": "State the Fineline's commercial role and what the evidence means for the product-development pipeline.",
        "opportunity_diagnosis": "Name the action or evidence gap in merchant and product-development language; avoid analytics jargon.",
        "product_proposal": "Frame a product direction or development brief, not an unsupported item promise.",
        "thank_you": "Keep the final page short, formal, and Paramont-branded.",
    }
    for slide_index, slide in enumerate(spec["slides"]):
        context = slide.get("fineline_context") or {}
        if context.get("white_space_coverage") == "missing" and slide["slide_role"] == "opportunity_diagnosis":
            allowed_types = ["validation_required"]
        elif slide["slide_role"] == "trend_visual":
            allowed_types = ["fact", "category_implication", "validation_required"]
        elif slide["slide_role"] == "product_proposal" and slide.get("limitations"):
            allowed_types = ["category_implication", "validation_required"]
        else:
            allowed_types = ["fact", "cross_analysis_synthesis", "category_implication", "proposal_recommendation", "validation_required"]
        locked_context = {
            "section": slide["section"], "slide_role": slide["slide_role"],
            "fineline": slide.get("fineline"), "purpose": slide["purpose"],
            "metric_cards": slide["metric_cards"], "chart": slide["chart"],
            "visual_assets": slide["visual_assets"], "limitations": slide["limitations"],
            "fineline_context": slide.get("fineline_context"),
            "product_fields": slide.get("product_fields"),
        }
        forbidden_phrases = []
        if slide_index < priority_index and slide["slide_role"] != "cover":
            forbidden_phrases = ["five priorities", "selected priorities", "selected finelines", "focus on five finelines"]
        briefs.append({
            "slide_id": slide["slide_id"],
            "communication_job": slide["purpose"],
            "audience": "Walmart category buyer",
            "locked_context": locked_context,
            "draft_narrative": {
                "headline": slide["headline"],
                "buyer_takeaway": slide["buyer_takeaway"],
                "insights": slide["insights"],
            },
            "allowed_evidence_ids": slide["source_refs"],
            "allowed_statement_types": allowed_types,
            "allowed_numeric_tokens": numeric_tokens({"locked_context": locked_context, "draft_narrative": {"headline": slide["headline"], "buyer_takeaway": slide["buyer_takeaway"]}}),
            "writing_rules": {
                "headline_mode": "fixed" if slide["slide_role"] in {"cover", "thank_you"} else "rewrite",
                "fixed_headline": slide["headline"] if slide["slide_role"] in {"cover", "thank_you"} else None,
                "headline_max_words": 18,
                "buyer_takeaway_max_words": 32,
                "insight_count_max": 3,
                "insight_max_words": 38,
                "voice": "concise, commercially fluent, buyer-facing, confident but evidence-bounded",
                "role_guidance": role_guidance.get(slide["slide_role"], "State the commercial implication in clear buyer-facing language."),
                "forbidden_phrases": forbidden_phrases,
                "avoid": ["analysis-process language", "generic overview titles", "repeated sentence templates", "unsupported certainty", "invented numbers"],
            },
        })
    return {
        "schema_version": "1.0",
        "task": "Rewrite only headline, buyer_takeaway, and insights for a Walmart buyer proposal deck.",
        "global_rules": [
            "Return JSON only and include every slide exactly once.",
            "Use only allowed evidence IDs and statement types for that slide.",
            "Do not alter or recompute locked facts, charts, rankings, layouts, visuals, or product fields.",
            "Do not introduce a numeric token absent from allowed_numeric_tokens.",
            "Write direct proposal language, not descriptions of what the slide should say.",
            "Build the story cumulatively: category context first, trends second, investment priorities third, then Fineline and product actions.",
            "Keep validation hypotheses visibly conditional and never promote them to facts.",
        ],
        "slides": briefs,
    }


def build(args: argparse.Namespace) -> tuple[dict[str, Any], dict[str, Any]]:
    artifact_descriptors: dict[str, dict[str, Any]] = {}
    if args.artifact_manifest:
        (fineline_data, fineline_descriptor), (ws_data, ws_descriptor) = load_manifest(
            Path(args.artifact_manifest), Path(args.workspace_root)
        )
        fineline_path = Path(fineline_descriptor.pop("_resolvedPath"))
        ws_path = Path(ws_descriptor.pop("_resolvedPath"))
        artifact_descriptors = {
            "fineline_investment_analysis": fineline_descriptor,
            "white_space_analysis": ws_descriptor,
        }
    else:
        if not args.fineline_json or not args.white_space_json:
            raise SystemExit("Use --artifact-manifest, or provide both --fineline-json and --white-space-json for local validation.")
        fineline_path, ws_path = map(Path, (args.fineline_json, args.white_space_json))
        fineline_data, ws_data = load_json(fineline_path), load_json(ws_path)
    if args.trend_mode == "enabled":
        if not args.trends_json or not args.moodboards_dir:
            raise SystemExit("--trend-mode enabled requires --trends-json and --moodboards-dir")
        trend_path = Path(args.trends_json)
        mood_dir: Path | None = Path(args.moodboards_dir)
        trends = load_json(trend_path)
    else:
        trend_path = fineline_path
        mood_dir = None
        trends = {"analysis_context": {"category": args.category}, "trends": []}
    paths = PathFormatter(args.path_mode)
    category = args.category
    discovered = [trends.get("analysis_context", {}).get("category"), fineline_data.get("summary", {}).get("category"), ws_data.get("analysis_context", {}).get("category")]
    mismatches = [x for x in discovered if x and category_key(x) != category_key(category)]
    if mismatches:
        raise SystemExit(f"Category mismatch: expected {category}; found {discovered}")
    if args.language != "en-US":
        raise SystemExit("Only --language en-US is currently supported")

    finelines = fineline_data.get("finelines") or []
    if not finelines:
        raise SystemExit("Fineline JSON contains no finelines")
    rows_by_name = {norm(x.get("final_fineline")): x for x in finelines}
    ws_by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in ws_data.get("facts") or []:
        ws_by_name[norm(row.get("final_fineline"))].append(row)
    investment_scores = calculate_investment_scores(finelines)
    eligible_names = [name for name, row in rows_by_name.items()
                      if str(row.get("opportunity_type")) in INVESTMENT_ELIGIBLE_TYPES]
    ranked_names = sorted(eligible_names, key=lambda n: investment_scores[n]["investment_score"], reverse=True)
    if args.top_n < 1 or args.top_n > len(ranked_names):
        raise SystemExit(f"--top-n must be between 1 and {len(ranked_names)} eligible Investment Finelines")
    auto_names = ranked_names[:args.top_n]
    user_names: list[str] = []
    unknown = []
    for requested in args.fineline or []:
        key = norm(requested)
        if key in rows_by_name and key not in user_names:
            user_names.append(key)
        elif key not in rows_by_name:
            unknown.append(requested)
    if unknown:
        names = [x.get("final_fineline") for x in finelines]
        details = {x: get_close_matches(x, names, n=3, cutoff=0.4) for x in unknown}
        raise SystemExit(f"Unknown explicit Finelines: {details}")
    selected_names = user_names + [name for name in auto_names if name not in user_names]
    selection_mode = "investment_plus_user_overrides"
    selected = [rows_by_name[x] for x in selected_names]
    investment_rank = {name: rank for rank, name in enumerate(ranked_names, 1)}
    selected_metadata: dict[str, dict[str, Any]] = {}
    for name in selected_names:
        is_user = name in user_names
        is_auto = name in auto_names
        coverage = "available" if ws_by_name.get(name) else "missing"
        selected_metadata[name] = {
            "selection_source": "investment_rule_and_user_override" if is_user and is_auto else "user_override" if is_user else "investment_rule",
            "investment_rank": investment_rank.get(name),
            "investment_opportunity_type": rows_by_name[name].get("opportunity_type"),
            "white_space_coverage": coverage,
            "proposal_readiness": "ready_for_diagnosis" if coverage == "available" else "white_space_validation_required",
        }

    ev = Evidence(paths)
    slides: list[dict[str, Any]] = []
    period = fineline_data.get("period", {})
    current_period = f"FY{str(period.get('current_period_start'))[:4]} WK{str(period.get('current_period_start'))[-2:]}–FY{str(period.get('current_period_end'))[:4]} WK{str(period.get('current_period_end'))[-2:]}"
    prior_period = f"FY{str(period.get('prior_period_start'))[:4]} WK{str(period.get('prior_period_start'))[-2:]}–FY{str(period.get('prior_period_end'))[:4]} WK{str(period.get('prior_period_end'))[-2:]}"
    total_sales = sum(finite(x.get("primary_current_sales_value")) for x in finelines)
    prior_sales = sum(finite(x.get("primary_prior_sales_value")) for x in finelines)
    category_growth = (total_sales / prior_sales - 1) if prior_sales else None
    total_units = sum(finite(x.get("primary_current_sales_units")) for x in finelines)
    prior_units = sum(finite(x.get("primary_prior_sales_units")) for x in finelines)
    category_unit_growth = (total_units / prior_units - 1) if prior_units else None
    category_asp = total_sales / total_units if total_units else None
    prior_category_asp = prior_sales / prior_units if prior_units else None
    category_asp_growth = (category_asp / prior_category_asp - 1) if prior_category_asp else None

    make_slide(slides, "opening", "cover", "Paramont Group - Walmart D19 Kids Craft",
               "Open the meeting with the business, retailer, department, and category context.",
               "Category Review and Product Development Proposal.",
               "light", layout="full_bleed_title")
    cat_cards = [
        card("Current Sales", total_sales, money(total_sales), "USD", current_period, prior_period, "Total across analyzed Finelines.", fineline_path, "$.finelines[*].primary_current_sales_value", ev),
        card("Category YoY", category_growth, pct(category_growth), "percent", current_period, prior_period, "WK31-aligned category change.", fineline_path, "$.finelines[*].primary_yoy_growth_pct", ev, category_growth),
        card("Current Units", total_units, f"{total_units / 1_000_000:.1f}M", "units", current_period, prior_period, "Total units across analyzed Finelines.", fineline_path, "$.finelines[*].primary_current_sales_units", ev),
        card("Average Selling Price", category_asp, f"${category_asp:.2f}", "USD per unit", current_period, prior_period, "Category weighted-average selling price.", fineline_path, "derived:category_sales/category_units", ev),
    ]
    make_slide(slides, "category", "category_snapshot", f"Kids Crafts Delivers {money(total_sales)} with {pct(category_growth)} Top-Line Growth",
               "Establish category scale, growth, unit demand, and average selling price before discussing where to invest.",
               "Start with how the total category is performing; Fineline priorities follow only after the category story is clear.",
               cards=cat_cards, insights=[insight("Category sales, units, and average selling price establish the commercial baseline for the product-development discussion.", "fact", [c["evidence_ids"][0] for c in cat_cards])], layout="metric_cards_statement")

    driver_rows = [
        {"metric": "sales", "current": total_sales, "prior": prior_sales, "yoy": category_growth},
        {"metric": "units", "current": total_units, "prior": prior_units, "yoy": category_unit_growth},
        {"metric": "asp", "current": category_asp, "prior": prior_category_asp, "yoy": category_asp_growth},
    ]
    driver_eid = ev.add(fineline_path, "derived:category_sales_units_asp", driver_rows, "fact", current_period, "mixed")
    make_slide(slides, "category", "category_growth_drivers", "Category Growth Reflects Both Demand and Price-Mix",
               "Explain total-category sales change through unit demand and average selling price before moving to Fineline differences.",
               "The balance of unit and ASP movement determines whether product development should prioritize demand creation, value architecture, or both.",
               chart={"chart_type": "grouped_bar", "business_question": "How did total Kids Crafts sales, units, and ASP change versus the aligned prior period?", "x_axis": "metric", "y_axis": "current_and_prior", "size": None, "color": "period", "series": ["current", "prior"], "sort": None, "filters": "category_total", "reference_lines": [], "highlight": None, "annotations": [], "data_rows": driver_rows, "source_file": paths.format(fineline_path), "source_json_paths": ["derived:category_sales_units_asp"], "evidence_ids": [driver_eid], "display": {"chart_title": "Category Sales, Units and ASP vs. Prior Year", "series_labels": {"current": "Current Period", "prior": "Prior Period"}, "metric_labels": {"sales": "Sales", "units": "Units", "asp": "Average Selling Price"}, "data_label_formats": {"sales": "$0.0M", "units": "0.0M", "asp": "$0.00"}, "data_label_policy": "show_current_and_prior", "axis_labels": {"x": "Category Metric", "y": "Use metric-specific formatted labels"}}},
               insights=[insight("Read category sales, unit demand, and average selling price together before assigning different roles to individual Finelines.", "category_implication", [driver_eid])], layout="chart_plus_insight")

    quadrant_specs = fineline_data.get("visualization_specs") or [{}]
    quad_spec = quadrant_specs[1] if len(quadrant_specs) > 1 else quadrant_specs[0]
    quad_rows = quad_spec.get("points") or [{"final_fineline": x["final_fineline"], "primary_yoy_growth_pct": x.get("primary_yoy_growth_pct"), "category_share": x.get("category_share"), "primary_current_sales_value": x.get("primary_current_sales_value"), "concentration_label": x.get("concentration_label")} for x in finelines]
    quad_source_path = "$.visualization_specs[1].points" if len(quadrant_specs) > 1 else "$.visualization_specs[0].points"
    quad_eid = ev.add(fineline_path, quad_source_path, quad_rows, "fact", current_period, "mixed")
    make_slide(slides, "category", "portfolio_quadrant", "Scale and Momentum Define Distinct Portfolio Roles",
               "Show the full category portfolio before narrowing to selected priorities.",
               "High-scale growth Finelines support innovation; high-scale slower spaces require defense and productivity actions.",
               chart={"chart_type": "bubble_quadrant", "business_question": "Which Finelines combine meaningful category scale with growth momentum?", "x_axis": quad_spec.get("x", "primary_yoy_growth_pct"), "y_axis": quad_spec.get("y", "category_share"), "size": quad_spec.get("size", "primary_current_sales_value"), "color": quad_spec.get("color", "concentration_label"), "series": [], "sort": None, "filters": quad_spec.get("chart_id", "zoomed_opportunity_quadrant"), "reference_lines": [{"axis": "x", "value": quad_spec.get("x_cutoff")}, {"axis": "y", "value": quad_spec.get("y_cutoff")}], "highlight": [x["final_fineline"] for x in selected], "annotations": [f"Largest share outlier excluded for readability: {quad_spec.get('excluded_outlier')}" if quad_spec.get("excluded_outlier") else "Focused comparison view"], "data_rows": quad_rows, "source_file": paths.format(fineline_path), "source_json_paths": ["$.visualization_specs[1].points" if len(quadrant_specs) > 1 else "$.visualization_specs[0].points"], "evidence_ids": [quad_eid], "display": {"chart_title": "Fineline Growth vs. Category Share", "axis_labels": {"x": "YoY Sales Growth", "y": "Category Share"}, "axis_number_formats": {"x": "0%", "y": "0.0%"}, "point_label_field": "final_fineline", "point_label_policy": "label_highlighted_plus_top_scale", "max_point_labels": 10, "bubble_size_legend": "Bubble size = Current Sales", "legend_labels": {"broad_based": "Broad-Based", "moderate_concentration": "Moderate Concentration", "high_concentration": "High Concentration"}}},
               insights=[insight("Portfolio role must reflect both relative scale and growth quality.", "category_implication", [quad_eid])], layout="full_width_chart")

    trend_rows = trends.get("trends") or []
    if args.trend_mode == "pending":
        trend_eid = ev.add(fineline_path, "workflow:trend_analysis_pending", "Trend analysis and mood boards are unavailable", "validation_required")
        make_slide(
            slides, "trends", "trend_pending", "Trend Evidence Will Be Added Before Concept Finalization",
            "Keep the trend chapter visible without inventing consumer or design evidence.",
            "Validate consumer shifts, design codes, and mood boards before finalizing product concepts.",
            "light", layout="validation_gap_statement",
            insights=[insight("Trend analysis and approved mood boards are not yet available.", "validation_required", [trend_eid])],
            limitations=["Trend analysis package and approved mood boards are pending."],
        )
    else:
        trend_overview = [{"trend_id": t.get("trend_id"), "trend_name": t.get("trend_name"), "evidence_strength": t.get("evidence_strength"), "confidence_score": t.get("confidence_score"), "trend_essence": t.get("trend_essence")} for t in trend_rows]
        trend_eid = ev.add(trend_path, "$.trends", trend_overview, "fact", None, None)
        make_slide(slides, "trends", "cultural_shift_overview", "Five Making Shifts Reframe What Kids Crafts Can Deliver" if len(trend_rows) == 5 else f"{len(trend_rows)} Making Shifts Reframe What Kids Crafts Can Deliver",
                   "Introduce the why-now consumer and design context.",
                   "Use trends to shape relevant experiences and design codes, not to claim Walmart sales demand.",
                   chart={"chart_type": "trend_cards", "business_question": "Which cross-report shifts matter to the category?", "x_axis": None, "y_axis": None, "size": "confidence_score", "color": "evidence_strength", "series": [], "sort": "trend_id", "filters": None, "reference_lines": [], "highlight": None, "annotations": [], "data_rows": trend_overview, "source_file": paths.format(trend_path), "source_json_paths": ["$.trends"], "evidence_ids": [trend_eid], "display": {"chart_title": "Five Consumer and Making Shifts Shaping Kids Crafts", "card_title_field": "trend_name", "card_body_field": "trend_essence", "card_meta_fields": ["evidence_strength", "confidence_score"], "card_body_max_words": 22, "data_label_policy": "show_title_body_and_strength"}},
                   insights=[insight("Treat these shifts as product and experience lenses; validate any commercial proposition against Walmart assortment data.", "validation_required", [trend_eid])], layout="trend_card_grid")

    mood_map = {p.name[:4].upper(): p for p in mood_dir.glob("*.png") if re.match(r"T\d{3}", p.name, re.I)} if mood_dir else {}
    for idx, trend in enumerate(trend_rows if args.trend_mode == "enabled" else []):
        tid = str(trend.get("trend_id") or "")
        essence_eid = ev.add(trend_path, f"$.trends[{idx}].trend_essence", trend.get("trend_essence"), "fact")
        why_eid = ev.add(trend_path, f"$.trends[{idx}].why_it_matters_for_category", trend.get("why_it_matters_for_category"), "category_implication")
        mood = mood_map.get(tid.upper())
        make_slide(slides, "trends", "trend_visual", f"{trend.get('trend_name')}: {trend.get('trend_essence')}",
                   "Translate one evidence-backed trend into buyer-relevant category direction.",
                   str(trend.get("why_it_matters_for_category") or "Use this trend as a product-design lens requiring assortment validation."),
                   "visual", layout="moodboard_hero", insights=[
                       insight(str(trend.get("trend_essence") or ""), "fact", [essence_eid]),
                       insight(str(trend.get("why_it_matters_for_category") or ""), "category_implication", [why_eid]),
                   ], visuals=[{"path": paths.format(mood) if mood else "MISSING", "usage": "Use as the full-width visual anchor; preserve aspect ratio and do not crop embedded evidence text.", "evidence_role": "visual_direction_only", "trend_id": tid}],
                   limitations=[] if mood else [f"Missing mood board for {tid}."])

    ranking_rows = [{
        "display_order": i + 1,
        "final_fineline": rows_by_name[n]["final_fineline"],
        **selected_metadata[n],
        **investment_scores[n],
    } for i, n in enumerate(selected_names)]
    ranking_eid = ev.add(fineline_path, "derived:investment_only_priority", ranking_rows, "cross_analysis_synthesis", current_period, "score")
    make_slide(slides, "priorities", "fineline_ranking", f"{len(selected)} Finelines Offer the Strongest Actionable Proposal Platforms",
               "Explain which Finelines receive detailed modules and why.",
               "Investment Analysis determines priority; White Space coverage determines readiness for deeper diagnosis, not selection.",
               chart={"chart_type": "horizontal_bar", "business_question": "Which Finelines should receive detailed proposal attention, and which are ready for White Space diagnosis?", "x_axis": "investment_score", "y_axis": "final_fineline", "size": None, "color": "proposal_readiness", "series": ["investment_score"], "sort": "display_order", "filters": selected_names, "reference_lines": [], "highlight": ranking_rows[0]["final_fineline"], "annotations": ["Show Investment rank, opportunity type, user override, White Space coverage, and proposal readiness"], "data_rows": ranking_rows, "source_file": paths.format(fineline_path), "source_json_paths": ["derived:investment_only_priority"], "evidence_ids": [ranking_eid], "display": {"chart_title": "Investment Priority and White Space Readiness", "axis_labels": {"x": "Investment Score", "y": "Fineline"}, "axis_number_formats": {"x": "0"}, "data_label_policy": "show_score_and_readiness", "category_label_field": "final_fineline", "secondary_label_fields": ["investment_opportunity_type", "white_space_coverage"]}},
               insights=[insight("White Space enriches selected Finelines but never changes Investment eligibility or rank.", "validation_required", [ranking_eid])], layout="ranking_bar_plus_readiness")

    products = load_products(Path(args.product_input) if args.product_input else None)
    products_by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for product in products:
        products_by_name[norm(product.get("final_fineline") or product.get("fineline"))].append(product)

    for fineline_index, row in enumerate(selected):
        name, key = str(row["final_fineline"]), norm(row["final_fineline"])
        base_path = f"$.finelines[{finelines.index(row)}]"
        cards = [
            card("Current Sales", row.get("primary_current_sales_value"), money(row.get("primary_current_sales_value")), "USD", current_period, prior_period, "Current Fineline scale.", fineline_path, f"{base_path}.primary_current_sales_value", ev),
            card("YoY Growth", row.get("primary_yoy_growth_pct"), pct(row.get("primary_yoy_growth_pct")), "percent", current_period, prior_period, str(row.get("growth_driver_label")), fineline_path, f"{base_path}.primary_yoy_growth_pct", ev, row.get("primary_yoy_growth_pct")),
            card("Category Share", row.get("category_share"), plain_pct(row.get("category_share")), "percent", current_period, "share of category", "Relative scale in category.", fineline_path, f"{base_path}.category_share", ev),
            card("Top-3 SKU Share", row.get("top3_sku_share"), plain_pct(row.get("top3_sku_share")), "percent", current_period, "share of Fineline sales", str(row.get("concentration_label")), fineline_path, f"{base_path}.top3_sku_share", ev),
        ]
        compare_rows = [{"metric": "sales", "current": row.get("primary_current_sales_value"), "prior": row.get("primary_prior_sales_value")}, {"metric": "units", "current": row.get("primary_current_sales_units"), "prior": row.get("primary_prior_sales_units")}, {"metric": "asp", "current": row.get("primary_current_asp"), "prior": row.get("primary_prior_asp")}]
        compare_eid = ev.add(fineline_path, f"{base_path}.{{primary_current_sales_value,primary_prior_sales_value,primary_current_sales_units,primary_prior_sales_units,primary_current_asp,primary_prior_asp}}", compare_rows, "fact", current_period, "mixed")
        make_slide(slides, "fineline_detail", "business_case", f"{name}: {str(row.get('opportunity_rationale') or 'Evidence Defines the Portfolio Role')}",
                   "Establish the Fineline business case using scale, momentum, and concentration.",
                   str(row.get("opportunity_rationale") or row.get("recommendation") or "Validate the Fineline role using item-level evidence."),
                   fineline=name, cards=cards,
                   chart={"chart_type": "grouped_bar", "business_question": "How did sales, units, and ASP change versus the aligned prior period?", "x_axis": "metric", "y_axis": "current_and_prior", "size": None, "color": "period", "series": ["current", "prior"], "sort": None, "filters": name, "reference_lines": [], "highlight": name, "annotations": [str(row.get("growth_driver_label"))], "data_rows": compare_rows, "source_file": paths.format(fineline_path), "source_json_paths": [f"{base_path}.primary_current_sales_value", f"{base_path}.primary_prior_sales_value", f"{base_path}.primary_current_sales_units", f"{base_path}.primary_prior_sales_units", f"{base_path}.primary_current_asp", f"{base_path}.primary_prior_asp"], "evidence_ids": [compare_eid], "display": {"chart_title": f"{name.title()} Sales, Units and ASP vs. Prior Year", "series_labels": {"current": "Current Period", "prior": "Prior Period"}, "metric_labels": {"sales": "Sales", "units": "Units", "asp": "Average Selling Price"}, "data_label_formats": {"sales": "$0.0M", "units": "0.0M", "asp": "$0.00"}, "data_label_policy": "show_current_and_prior", "axis_labels": {"x": "Metric", "y": "Use metric-specific formatted labels"}}},
                   insights=[
                       insight(f"{name} delivers {money(row.get('primary_current_sales_value'))} in current-period sales with {pct(row.get('primary_yoy_growth_pct'))} YoY change.", "fact", [cards[0]["evidence_ids"][0], cards[1]["evidence_ids"][0]]),
                       insight(f"Top-three SKUs represent {plain_pct(row.get('top3_sku_share'))} of sales, classified as {row.get('concentration_label')}.", "fact", [cards[3]["evidence_ids"][0]]),
                       insight(str(row.get("opportunity_rationale") or "Item-level validation is required."), "proposal_recommendation", [compare_eid]),
                   ], layout="metric_cards_plus_driver_chart")

        ws_rows = best_ws_rows(ws_by_name.get(key, []))
        ws_data_rows = [{k: ws.get(k) for k in ("final_fineline", "opportunity_id", "attribute_combination", "strategy_code", "current_sales_value", "sales_yoy_growth", "recent_8w_productivity", "mean_stores_selling", "median_sales_per_store_week", "tag_coverage_rate", "recommended_action", "risk_flags")} for ws in ws_rows]
        ws_eid = ev.add(ws_path, f"$.facts[?(@.final_fineline=='{name}')]", ws_data_rows, "fact", "recent 8 weeks and analysis period", "mixed")
        positive = [x for x in ws_rows if POSITIVE_WS.get(str(x.get("strategy_code")), 0.5) >= 0.65]
        best = ws_rows[0] if ws_rows else None
        ws_cards = []
        if best:
            candidates = [
                ("White-space Sales", "current_sales_value", money, "USD", "analysis period", "selected attribute space", str(best.get("attribute_combination"))),
                ("White-space YoY", "sales_yoy_growth", pct, "percent", "analysis period", "prior period", str(best.get("strategy_code"))),
                ("Recent Productivity", "recent_8w_productivity", lambda v: f"{finite(v):,.1f}", "sales productivity", "recent 8 weeks", "attribute opportunity", "Directional productivity signal."),
                ("Tag Coverage", "tag_coverage_rate", plain_pct, "percent", "analysis period", "tagged rows", "Evidence coverage for the attribute conclusion."),
            ]
            for label, field, formatter, unit, card_period, comparison, interpretation_text in candidates:
                if best.get(field) is not None:
                    ws_cards.append(card(label, best.get(field), formatter(best.get(field)), unit, card_period, comparison, interpretation_text, ws_path, f"$.facts[?(@.opportunity_id=='{best.get('opportunity_id')}')].{field}", ev, best.get(field)))
        ws_insights = (
            [
                insight(str(best.get("why_this_action")), "proposal_recommendation", [ws_eid]),
                insight(f"{len(positive)} positive action signal(s) are available among the displayed attribute spaces.", "cross_analysis_synthesis", [ws_eid]),
                insight("Candidate actions require item, cost, price, and operational validation before buyer commitment.", "validation_required", [ws_eid]),
            ] if best else [
                insight("White-space evidence is missing for this Fineline; run White Space Analysis before making an attribute, distribution, or entry recommendation.", "validation_required", [ws_eid]),
                insight("Investment priority is retained, but proposal readiness is limited to the Business Case and trend-led hypothesis.", "validation_required", [ws_eid]),
            ]
        )
        make_slide(slides, "fineline_detail", "opportunity_diagnosis", f"{name}: White Space Points to {str(best.get('strategy_code')).replace('_', ' ').title() if best else 'Further Validation'}",
                   "Translate attribute and distribution evidence into a bounded action.",
                   str(best.get("recommended_action")) if best else "No white-space fact is available; obtain attribute-level evidence before making a product claim.",
                   fineline=name, cards=ws_cards,
                   chart={"chart_type": "scatter", "business_question": "Which attribute spaces combine growth and recent productivity?", "x_axis": "sales_yoy_growth", "y_axis": "recent_8w_productivity", "size": "current_sales_value", "color": "strategy_code", "series": [], "sort": None, "filters": name, "reference_lines": [{"axis": "x", "value": 0}], "highlight": best.get("opportunity_id"), "annotations": [x.get("attribute_combination") for x in ws_rows[:3]], "data_rows": ws_data_rows, "source_file": paths.format(ws_path), "source_json_paths": [f"$.facts[?(@.final_fineline=='{name}')]"], "evidence_ids": [ws_eid], "display": {"chart_title": "Attribute Growth vs. Recent Productivity", "axis_labels": {"x": "YoY Sales Growth", "y": "Recent 8-Week Sales Productivity"}, "axis_number_formats": {"x": "0%", "y": "0"}, "point_label_field": "attribute_combination", "point_label_policy": "label_highlighted_plus_top_opportunities", "max_point_labels": 6, "label_cleanup": "replace underscores with spaces; remove field-name prefix when space permits", "bubble_size_legend": "Bubble size = Current Sales", "legend_field": "strategy_code"}} if best else None,
                   insights=ws_insights,
                   limitations=[] if best else ["No matching white-space rows; do not present a specific White Space conclusion."],
                   layout="metric_cards_plus_scatter" if best else "validation_gap_statement")

        product_rows = products_by_name.get(key, [])
        images = image_matches(Path(args.product_images_dir) if args.product_images_dir else None, name, paths)
        if product_rows:
            fields = dict(product_rows[0])
            for field in TBD_FIELDS:
                if fields.get(field) in (None, ""):
                    fields[field] = "TBD"
            product_source = Path(args.product_input)
            product_eid = ev.add(product_source, f"records[final_fineline='{name}']", product_rows, "fact")
            title = f"{name}: Translate the Opportunity into a Buyer-Ready Product Proposal"
            takeaway = "Use supplied product facts and retain TBD for every unsupported commercial claim."
            product_insights = [insight("Product details are supplied inputs and require normal commercial review.", "validation_required", [product_eid])]
        else:
            fields = {field: "TBD" for field in TBD_FIELDS}
            fields.update({"consumer_need": "Derive from the selected trend lens and Fineline evidence; validation required.", "trend_alignment": [t.get("trend_id") for t in trend_rows], "suggested_product_format": "TBD — define after concept development.", "play_or_making_mechanic": "TBD — define after concept development.", "design_language": "Use applicable mood boards as visual direction only.", "shelf_role": "TBD", "validation_hypothesis": "A differentiated concept can address the documented Fineline opportunity."})
            if args.trend_mode == "pending":
                fields.update({
                    "consumer_need": "PENDING — define after trend validation.",
                    "trend_alignment": "PENDING",
                    "design_language": "PENDING — validate with trend analysis and approved mood boards.",
                })
            product_eid = ev.add(fineline_path, f"{base_path}.opportunity_type", row.get("opportunity_type"), "category_implication", current_period)
            title = f"{name}: Convert the Evidence into a Product Direction"
            takeaway = "The opportunity is evidence-backed; product, price, economics, and execution remain to be developed."
            product_insights = [insight("Do not treat a trend-aligned direction as a validated item opportunity.", "validation_required", [product_eid])]
        product_visuals = [{"path": p, "usage": "Primary product visual", "evidence_role": "supplied_product_asset"} for p in images]
        make_slide(slides, "fineline_detail", "product_proposal", title,
                   "Provide the structured product handoff immediately after the Fineline diagnosis.", takeaway,
                   "visual", fineline=name, layout="product_hero_with_fact_panel", insights=product_insights,
                   visuals=product_visuals, product_fields=fields,
                   limitations=([] if product_rows else ["No product input supplied; all commercial product fields remain TBD."])
                   + (["Trend analysis and approved mood boards are pending."] if args.trend_mode == "pending" else []))

    action_counts = defaultdict(int)
    for row in selected:
        action_counts[str(row.get("opportunity_type") or "unclassified")] += 1
    portfolio_rows = [{"opportunity_type": k, "fineline_count": v, "finelines": [r["final_fineline"] for r in selected if str(r.get("opportunity_type") or "unclassified") == k]} for k, v in sorted(action_counts.items())]
    portfolio_eid = ev.add(fineline_path, "derived:selected opportunity_type counts", portfolio_rows, "cross_analysis_synthesis", current_period, "count")
    make_slide(slides, "closing", "portfolio_recommendation", "A Balanced Action Portfolio Matches Investment to Evidence",
               "Summarize the proposed mix of innovation, expansion, refresh, validation, and rationalization.",
               "Advance the strongest actions while keeping validation gates explicit.",
               chart={"chart_type": "portfolio_action_matrix", "business_question": "How should the selected Finelines be managed as a portfolio?", "x_axis": "opportunity_type", "y_axis": "fineline_count", "size": None, "color": "opportunity_type", "series": [], "sort": "fineline_count_desc", "filters": selected_names, "reference_lines": [], "highlight": None, "annotations": [], "data_rows": portfolio_rows, "source_file": paths.format(fineline_path), "source_json_paths": ["derived:selected opportunity_type counts"], "evidence_ids": [portfolio_eid], "display": {"chart_title": "Recommended Portfolio Roles", "axis_labels": {"x": "Number of Finelines", "y": "Portfolio Role"}, "data_label_policy": "show_count_and_fineline_names", "category_label_field": "opportunity_type", "category_label_cleanup": "replace underscores with spaces"}},
               insights=[insight("Stage actions by evidence readiness instead of presenting every Fineline as a new-item request.", "proposal_recommendation", [portfolio_eid])], layout="portfolio_matrix")
    make_slide(slides, "closing", "buyer_ask", "Align on Priority Platforms, Validation Gates, and the Next Review",
               "Convert analysis into a bounded buyer decision and workplan.",
               "Agree which Finelines advance to concept review, which need distribution or productivity work, and what evidence is required next.",
               "light", layout="three_step_ask", insights=[
                   insight("Confirm the priority Finelines and intended portfolio role.", "proposal_recommendation", [ranking_eid]),
                   insight("Agree the product, price, margin, item, and store-level evidence required for concept approval.", "validation_required", [ranking_eid]),
                   insight("Set the timing and owners for the next concept and assortment review.", "proposal_recommendation", [ranking_eid]),
               ])
    make_slide(slides, "closing", "method_notes", "Every Proposal Claim Remains Traceable to Its Source",
               "Explain scoring, evidence boundaries, periods, and limitations.",
               "The proposal is a decision framework, not a substitute for item economics, assortment review, or buyer approval.",
               "light", layout="method_notes", insights=[
                   insight("Trend evidence explains why-now and design direction; it does not prove Walmart sales demand.", "validation_required", [trend_eid]),
                   insight("Fineline and white-space inputs retain their original periods, units, and data-quality warnings.", "fact", [ranking_eid, ws_eid]),
               ])
    make_slide(slides, "closing", "thank_you", "Thank You",
               "Close the buyer proposal with a clean Paramont-branded final page.",
               "Paramont Group appreciates the opportunity to support Walmart D19 Kids Craft growth.",
               "light", layout="thank_you", insights=[])

    for slide in slides:
        if slide.get("fineline"):
            slide["fineline_context"] = selected_metadata[norm(slide["fineline"])]

    spec = {
        "schema_version": "1.0", "generated_at": datetime.now(timezone.utc).isoformat(),
        "deck_title": f"Walmart {category.title()} Growth Proposal", "category": category,
        "language": args.language,
        "inputs": {"trend_mode": args.trend_mode, "trends_json": paths.format(trend_path) if args.trend_mode == "enabled" else None, "moodboards_dir": paths.format(mood_dir), "fineline_json": paths.format(fineline_path), "white_space_json": paths.format(ws_path), "product_input": paths.format(Path(args.product_input)) if args.product_input else None, "product_images_dir": paths.format(Path(args.product_images_dir)) if args.product_images_dir else None},
        "selection": {
            "mode": selection_mode,
            "auto_top_n": args.top_n,
            "investment_eligible_types": sorted(INVESTMENT_ELIGIBLE_TYPES),
            "user_overrides": [rows_by_name[n]["final_fineline"] for n in user_names],
            "selected_finelines": [x["final_fineline"] for x in selected],
            "ranking_basis": "investment_only",
            "ranking": ranking_rows,
            "white_space_declared_finelines": ws_data.get("analysis_context", {}).get("finelines") or [],
        },
        "periods": {"current": current_period, "prior": prior_period},
        "slides": slides,
        "product_tbd_fields": TBD_FIELDS,
        "narrative_status": "draft_template",
        "ppt_execution_contract": {
            "chart_labels_required": True,
            "rules": [
                "Use chart.display.chart_title as the audience-facing chart title.",
                "Use chart.display.axis_labels; never expose raw JSON field names on axes.",
                "Apply chart.display data-label formats and point-label policy.",
                "Use business-readable legend labels and remove redundant legends.",
                "Format sales in $M or $K, units in M or K, ASP as currency, and rates as percentages.",
                "Every scatter or bubble point shown to the buyer must be identifiable according to point_label_policy.",
                "Shorten or selectively label crowded charts rather than showing overlapping or unreadable text.",
            ],
        },
    }
    fineline_source = paths.format(fineline_path)
    white_space_source = paths.format(ws_path)
    for record in ev.records:
        if record.get("source_file") == fineline_source:
            record["source_layer"] = "fineline_investment_analysis"
        elif record.get("source_file") == white_space_source:
            record["source_layer"] = "white_space_analysis"
    evidence_map = {"schema_version": "1.0", "generated_at": spec["generated_at"], "records": ev.records}
    spec["_artifact_descriptors"] = artifact_descriptors
    return spec, evidence_map


def to_presentation_outline(spec: dict[str, Any], evidence_map: dict[str, Any]) -> dict[str, Any]:
    """Project the completed domain plan directly into the Bento handoff contract."""
    descriptors = spec.pop("_artifact_descriptors", {})
    source_by_layer: dict[str, str] = {}
    sources: list[dict[str, Any]] = []
    for layer, descriptor in descriptors.items():
        source_id = f"SRC{len(sources) + 1:03d}"
        source_by_layer[layer] = source_id
        sources.append({"sourceId": source_id, **descriptor})
    for record in evidence_map["records"]:
        layer = record.get("source_layer") or "other"
        if layer not in source_by_layer:
            source_id = f"SRC{len(sources) + 1:03d}"
            source_by_layer[layer] = source_id
            sources.append({
                "sourceId": source_id,
                "producerSkill": layer,
                "relativePath": record.get("source_file"),
            })
    statement_map = {
        "cross_analysis_synthesis": "cross_source_synthesis",
        "category_implication": "implication",
        "proposal_recommendation": "recommendation",
    }
    evidence = [{
        "evidenceId": row["evidence_id"],
        "sourceId": source_by_layer[row.get("source_layer") or "other"],
        "fieldPath": row.get("source_json_path"),
        "value": row.get("value"),
        "period": row.get("period"),
        "unit": row.get("unit"),
        "statementType": statement_map.get(row.get("statement_type"), row.get("statement_type")),
        "note": row.get("note"),
    } for row in evidence_map["records"]]

    slides: list[dict[str, Any]] = []
    for old in spec["slides"]:
        blocks: list[dict[str, Any]] = []
        if old["metric_cards"]:
            blocks.append({
                "id": f"{old['slide_id']}-kpis", "type": "kpi_group", "role": "primary",
                "items": [{
                    "id": f"{old['slide_id']}-kpi-{index + 1}",
                    "label": item["label"], "rawValue": item["raw_value"],
                    "displayValue": item["display_value"], "unit": item["unit"],
                    "period": item["period"], "comparisonBasis": item["comparison_basis"],
                    "direction": item["direction"], "interpretation": item["interpretation"],
                    "evidenceIds": item["evidence_ids"],
                } for index, item in enumerate(old["metric_cards"])],
                "evidenceIds": sorted({eid for item in old["metric_cards"] for eid in item["evidence_ids"]}),
            })
        if old.get("chart"):
            chart = old["chart"]
            blocks.append({
                "id": f"{old['slide_id']}-chart", "type": "chart", "role": "primary",
                "chart": {
                    "type": chart["chart_type"], "businessQuestion": chart["business_question"],
                    "data": {"rows": chart["data_rows"]},
                    "encoding": {"x": chart.get("x_axis"), "y": chart.get("y_axis"),
                                 "size": chart.get("size"), "color": chart.get("color"),
                                 "series": chart.get("series")},
                    "display": chart.get("display") or {},
                    "referenceLines": chart.get("reference_lines") or [],
                    "highlights": chart.get("highlight") if isinstance(chart.get("highlight"), list) else ([chart["highlight"]] if chart.get("highlight") else []),
                    "annotations": chart.get("annotations") or [],
                },
                "evidenceIds": chart.get("evidence_ids") or [],
            })
        for index, visual in enumerate(old["visual_assets"]):
            blocks.append({
                "id": f"{old['slide_id']}-image-{index + 1}", "type": "image", "role": "primary",
                "source": {"relativePath": visual.get("path")}, "usage": visual.get("usage"),
                "fit": "contain", "evidenceIds": [],
            })
        if old["insights"]:
            blocks.append({
                "id": f"{old['slide_id']}-insights", "type": "insight_list", "role": "supporting",
                "items": [{
                    "text": item["text"],
                    "statementType": statement_map.get(item["statement_type"], item["statement_type"]),
                    "evidenceIds": item["evidence_ids"],
                } for item in old["insights"]],
                "evidenceIds": sorted({eid for item in old["insights"] for eid in item["evidence_ids"]}),
            })
        if old.get("product_fields") is not None:
            blocks.append({
                "id": f"{old['slide_id']}-facts", "type": "fact_panel", "role": "supporting",
                "fields": [{"label": key, "value": value, "status": "tbd" if value == "TBD" else "ready"}
                           for key, value in old["product_fields"].items()],
                "evidenceIds": old["source_refs"],
            })
        if old["slide_role"] == "trend_pending":
            blocks.insert(0, {
                "id": f"{old['slide_id']}-validation-gap", "type": "validation_gap", "role": "primary",
                "message": old["limitations"][0],
                "requiredEvidence": ["Trend analysis package", "Approved mood boards"],
                "evidenceIds": old["source_refs"],
            })
        if not blocks:
            blocks.append({"id": f"{old['slide_id']}-text", "type": "text", "role": "primary", "text": old["buyer_takeaway"], "evidenceIds": old["source_refs"]})
        extension = {}
        if old.get("fineline") or old.get("fineline_context"):
            extension = {"walmartBuyerProposal": {"finalFineline": old.get("fineline"), **(old.get("fineline_context") or {})}}
        slides.append({
            "slideId": old["slide_id"], "section": old["section"], "role": old["slide_role"],
            "objective": old["purpose"], "audienceDecision": old["buyer_takeaway"],
            "headline": old["headline"], "takeaway": old["buyer_takeaway"],
            "layoutHint": old["layout"], "density": old["density"], "blocks": blocks,
            "speakerNotes": {"talkTrack": old["buyer_takeaway"], "evidenceIds": old["source_refs"]},
            "limitations": old["limitations"], "extensions": extension,
        })
    return {
        "contract": "paramont.presentation-outline/v1",
        "generatedAt": spec["generated_at"],
        "producer": {"skillName": "build-walmart-buyer-proposal-outline", "operation": "buildOutline"},
        "document": {"id": slug(spec["deck_title"]), "title": spec["deck_title"], "language": spec["language"],
                     "audience": "Walmart category buyer", "purpose": "Support assortment and product-development decisions"},
        "theme": {"templateId": "wmt-kids-mod", "aspectRatio": "16:9"},
        "sources": sources, "evidence": evidence, "slides": slides,
        "narrative": {"status": "draft", "source": None, "appliedAt": None},
        "validation": {"status": "draft", "errors": [], "warnings": []},
        "extensions": {"walmartBuyerProposal": {
            "category": spec["category"], "periods": spec.get("periods"), "selection": spec["selection"],
            "productTbdFields": spec["product_tbd_fields"], "inputs": spec["inputs"],
        }},
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--artifact-manifest")
    parser.add_argument("--workspace-root", default=".")
    parser.add_argument("--trend-mode", choices=("enabled", "pending"), default="enabled")
    parser.add_argument("--trends-json")
    parser.add_argument("--moodboards-dir")
    parser.add_argument("--fineline-json")
    parser.add_argument("--white-space-json")
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--category", default="KIDS CRAFTS")
    parser.add_argument("--top-n", type=int, default=10)
    parser.add_argument("--fineline", action="append")
    parser.add_argument("--language", default="en-US")
    parser.add_argument("--product-input")
    parser.add_argument("--product-images-dir")
    parser.add_argument("--path-mode", choices=("relative", "absolute"), default="relative",
                        help="Write source and visual paths as relative paths by default; use absolute for local debugging.")
    parser.add_argument(
        "--narrative-mode",
        choices=("external", "draft"),
        default="draft",
        help=(
            "draft: write proposal_narrative.json from brief draft_narrative and mark "
            "outline narrative as llm_enriched (deterministic chain; no hand-written copy). "
            "external: leave narrative.status=draft for a later apply_narrative.py step."
        ),
    )
    args = parser.parse_args()
    spec, evidence_map = build(args)
    narrative_brief = build_narrative_brief(spec)
    markdown = render_markdown(spec)
    outline = to_presentation_outline(spec, evidence_map)
    output = Path(args.output_dir)
    output.mkdir(parents=True, exist_ok=True)
    if str(getattr(args, "narrative_mode", "external") or "external") == "draft":
        narrative_doc = {
            "schema_version": "1.0",
            "slides": [
                {
                    "slide_id": slide["slide_id"],
                    "headline": (slide.get("draft_narrative") or {}).get("headline")
                    or "",
                    "buyer_takeaway": (slide.get("draft_narrative") or {}).get(
                        "buyer_takeaway"
                    )
                    or "",
                    "insights": list(
                        (slide.get("draft_narrative") or {}).get("insights") or []
                    ),
                }
                for slide in narrative_brief.get("slides") or []
            ],
        }
        (output / "proposal_narrative.json").write_text(
            json.dumps(narrative_doc, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        outline["narrative"] = {
            "status": "llm_enriched",
            "source": "draft",
            "appliedAt": datetime.now(timezone.utc).isoformat(),
        }
    (output / "proposal_presentation_outline.json").write_text(json.dumps(outline, ensure_ascii=False, indent=2), encoding="utf-8")
    (output / "proposal_narrative_brief.json").write_text(json.dumps(narrative_brief, ensure_ascii=False, indent=2), encoding="utf-8")
    (output / "proposal_deck_outline.md").write_text(markdown, encoding="utf-8")
    print(json.dumps({"status": "built", "slides": len(outline["slides"]), "selected_finelines": outline["extensions"]["walmartBuyerProposal"]["selection"]["selected_finelines"], "output_dir": str(output.resolve())}, ensure_ascii=False))


if __name__ == "__main__":
    main()
